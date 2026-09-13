"""
Script to extract numeric ratings from makerspace study Excel files and append them as new columns.
Usage:
  python makerspace_preprocessing2.py                         # uses defaults below
Run  python makerspace_preprocessing2.py --help  for all options.
Author: Mariam Sulleiman 23/04/2026
Preprocesses all Excel (.xlsx) files found in a given input directory and
applies one transformation:
"""

import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

# DEFAULT CONFIGURATION

# Directory that contains the Excel files to process
DEFAULT_INPUT_DIR = "C:\\Users\\Mariam Sulleiman\\LLM\\API_data"

# Directory where processed files are saved (created automatically if absent)
DEFAULT_OUTPUT_DIR = "processed_output"

# Name of the sheet targeted for rating extraction.
# Set to None (or pass --sheet "" on the CLI) to always use the last sheet.
DEFAULT_TARGET_SHEET = "All_Data"

# Row index (0-based) that contains the persona names (used to build rating headers)
PERSONA_ROW_INDEX = 1  # Row 2 in Excel = index 1 in Python

# Suffix appended to each persona name to form the rating column header
RATING_HEADER_SUFFIX = "_Rating"

# Output filename suffix (inserted before the file extension)
OUTPUT_SUFFIX = "_processed"

#  RATING EXTRACTION HELPERS

def extract_leading_number(value):
    """
    Return the first integer found at the START of `value` as an int,
    or None if the cell does not begin with a digit.

    Examples
    --------
    '4\\nJustification: ...'   ->  4
    '2-3 (Minor to Moderate)' ->  2
    'Dangerous-Safe: 4\\n...'  ->  None  (does not start with a digit)
    'Making a wooden pen.'     ->  None
    """
    if value is None:
        return None
    text = str(value).strip()
    match = re.match(r"^(\d+)", text)
    return int(match.group(1)) if match else None


def add_rating_columns(ws):
    """
    Append rating columns to the RIGHT of the existing data columns.

    For each existing data column (col B onward), a new rating column is
    added at the end of the sheet. Every row gets the extracted leading
    integer from the corresponding data cell placed in its rating column.

    The persona row (PERSONA_ROW_INDEX) receives a header like
    'Neutral_Rating', 'Woman_Rating', 'Man_Rating'.

    Returns the number of rating columns added.
    """
    # Read all rows into memory (we will write back column by column)
    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        return 0

    n_data_cols = len(all_rows[0])   # total columns including col A (label col)
    n_personas  = n_data_cols - 1    # columns B onward = one per persona

    # Persona names come from PERSONA_ROW_INDEX, columns B onward (index 1+)
    persona_row = all_rows[PERSONA_ROW_INDEX]
    persona_names = [str(persona_row[i]) if persona_row[i] else f"Col{i}" 
                     for i in range(1, n_data_cols)]

    # First empty column index (1-based) for the new rating columns
    first_rating_col = n_data_cols + 1  # openpyxl columns are 1-based

    for persona_offset, persona_name in enumerate(persona_names):
        data_col_index  = persona_offset + 1   # 0-based index into each row tuple
        rating_col      = first_rating_col + persona_offset  # 1-based Excel column

        for row_index, row in enumerate(all_rows):
            excel_row = row_index + 1  # convert to 1-based

            if row_index == PERSONA_ROW_INDEX:
                # Write the rating column header in the persona row
                header = f"{persona_name}{RATING_HEADER_SUFFIX}"
                ws.cell(row=excel_row, column=rating_col, value=header)
            else:
                # Extract leading number from the data cell and write rating
                cell_value = row[data_col_index] if data_col_index < len(row) else None
                rating = extract_leading_number(cell_value)
                ws.cell(row=excel_row, column=rating_col, value=rating)

    return n_personas

#  PIPELINE

def process_file(input_path, output_path, target_sheet):
    """
    Load workbook, append rating columns in the target sheet, and save.
    """
    print(f"\n  Processing : {input_path.name}")

    wb = load_workbook(input_path)

    # Locate target sheet
    if target_sheet and target_sheet in wb.sheetnames:
        ws_target = wb[target_sheet]
        print(f"  Target sheet : '{target_sheet}'")
    else:
        ws_target = wb[wb.sheetnames[-1]]
        print(f"  Target sheet : '{ws_target.title}' (last sheet — '{target_sheet}' not found)")

    n_cols = add_rating_columns(ws_target)
    print(f"  Rating columns added : {n_cols}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"  Saved -> {output_path}")


#  VERIFICATION HELPER

def verify_output(output_path, target_sheet):
    """Print a preview showing original data columns and new rating columns."""
    wb = load_workbook(output_path)

    if target_sheet and target_sheet in wb.sheetnames:
        ws = wb[target_sheet]
    else:
        ws = wb[wb.sheetnames[-1]]

    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        print("  (sheet is empty)")
        return

    n_cols = len(all_rows[0])

    print(f"\n  Sheet '{ws.title}' — {ws.max_row} rows x {n_cols} cols")
    print(f"  {'Row Label':<30} | {'Data cols (B-D)':<45} | Rating cols (E onward)")
    print("  " + "-" * 110)

    for row in all_rows:
        label     = str(row[0])[:28] if row[0] else ""
        data_vals = "  |  ".join(str(v)[:12] if v is not None else "None" for v in row[1:4])
        rate_vals = "  |  ".join(str(v)[:12] if v is not None else "None" for v in row[4:])
        print(f"  {label:<30} | {data_vals:<45} | {rate_vals}")


#  CLI & ENTRY POINT

def parse_args():
    parser = argparse.ArgumentParser(
        description=(
            "Preprocess makerspace study Excel files:\n"
            "  Append numeric rating columns side-by-side with persona columns\n"
            "  in the All_Data sheet."
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument(
        "--input_dir",
        default=DEFAULT_INPUT_DIR,
        help=f"Directory containing input .xlsx files (default: '{DEFAULT_INPUT_DIR}')",
    )
    parser.add_argument(
        "--output_dir",
        default=DEFAULT_OUTPUT_DIR,
        help=f"Directory for processed output files (default: '{DEFAULT_OUTPUT_DIR}')",
    )
    parser.add_argument(
        "--sheet",
        default=DEFAULT_TARGET_SHEET,
        help=(
            f"Sheet to use for rating extraction (default: '{DEFAULT_TARGET_SHEET}'). "
            "Pass an empty string to always use the last sheet."
        ),
    )
    parser.add_argument(
        "--verify",
        action="store_true",
        help="Print a preview of the output sheet after processing.",
    )
    return parser.parse_args()


def main():
    args = parse_args()

    input_dir    = Path(args.input_dir)
    output_dir   = Path(args.output_dir)
    target_sheet = args.sheet or None

    # Validate input directory
    if not input_dir.is_dir():
        print(f"ERROR: Input directory not found: {input_dir}", file=sys.stderr)
        sys.exit(1)

    # Discover Excel files
    input_files = sorted(input_dir.glob("*.xlsx"))
    if not input_files:
        print(f"No .xlsx files found in '{input_dir}'. Nothing to do.")
        sys.exit(0)

    print(f"Found {len(input_files)} file(s) in '{input_dir}':")
    for f in input_files:
        print(f"  * {f.name}")

    # Process each file
    output_dir.mkdir(parents=True, exist_ok=True)
    results = []

    for src in input_files:
        out_name = src.stem + OUTPUT_SUFFIX + src.suffix
        out_path = output_dir / out_name
        try:
            process_file(src, out_path, target_sheet)
            results.append((src, out_path, None))
        except Exception as exc:
            print(f"  ERROR processing {src.name}: {exc}", file=sys.stderr)
            results.append((src, out_path, exc))

    # Optional verification
    if args.verify:
        print("\n" + "=" * 70)
        print("VERIFICATION REPORT")
        print("=" * 70)
        for src, out_path, err in results:
            if err:
                print(f"\n  [SKIPPED - error] {src.name}")
                continue
            print(f"\n  File: {out_path.name}")
            verify_output(out_path, target_sheet)

    # Summary
    print("\n" + "=" * 70)
    ok_count  = sum(1 for _, _, e in results if e is None)
    err_count = sum(1 for _, _, e in results if e is not None)
    print(f"Done. {ok_count} file(s) processed successfully", end="")
    print(f", {err_count} error(s)." if err_count else ".")
    print(f"Output directory: {output_dir.resolve()}")

    if err_count:
        sys.exit(1)


if __name__ == "__main__":
    main()
